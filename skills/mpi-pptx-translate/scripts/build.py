# /// script
# requires-python = ">=3.9"
# dependencies = [
#   "python-pptx",
#   "pyyaml",
# ]
# ///

import sys, yaml
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import MSO_AUTO_SIZE

FONT_SCALE = 0.82  # shrink ~18%


def shrink_font_tf(tf):
    for para in tf.paragraphs:
        for run in para.runs:
            if run.font.size:
                run.font.size = Pt(int(run.font.size.pt * FONT_SCALE))
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass


def build(yaml_path, src_path, out_path):
    with open(yaml_path) as f:
        entries = yaml.safe_load(f)

    index = {}
    for e in entries:
        index[(e["slide"], e["shape"], e["run"])] = e["en"]

    prs = Presentation(src_path)

    for slide_num, slide in enumerate(prs.slides, 1):
        for shape_idx, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                has_translation = False
                for para_idx, para in enumerate(shape.text_frame.paragraphs):
                    key = (slide_num, shape_idx, para_idx)
                    if key in index:
                        has_translation = True
                        en = index[key]
                        for r in para.runs:
                            r.text = ""
                        if para.runs:
                            para.runs[0].text = en
                        else:
                            para.add_run().text = en
                if has_translation:
                    shrink_font_tf(shape.text_frame)

            elif shape.has_table:
                num_rows = len(shape.table.rows)
                for r in range(num_rows * len(shape.table.columns)):
                    key = (slide_num, shape_idx, r)
                    if key in index:
                        row = r % num_rows
                        col = r // num_rows
                        shape.table.cell(row, col).text = index[key]
                for row in shape.table.rows:
                    for cell in row.cells:
                        shrink_font_tf(cell.text_frame)

        key = (slide_num, -1, 0)
        if key in index and slide.has_notes_slide:
            ns = slide.notes_slide
            ns.notes_text_frame.clear()
            ns.notes_text_frame.paragraphs[0].add_run().text = index[key]

    prs.save(out_path)
    print(f"Saved {out_path}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("Usage: uv run build.py <strings.yaml> <input.pptx> <output.pptx>", file=sys.stderr)
        sys.exit(1)
    build(sys.argv[1], sys.argv[2], sys.argv[3])
