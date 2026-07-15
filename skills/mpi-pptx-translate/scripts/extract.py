#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.9"
# dependencies = [
#   "python-pptx",
#   "pyyaml",
# ]
# ///

import sys, yaml
from pptx import Presentation

PLACEHOLDER_KINDS = {
    1: "title", 2: "body", 3: "center_title",
    4: "subtitle", 5: "body", 6: "body", 7: "body",
}


def shape_kind(shape):
    if shape.has_table:
        return "table"
    try:
        ph = shape.placeholder_format
        if ph is not None and ph.type is not None:
            return PLACEHOLDER_KINDS.get(ph.type, "body")
    except ValueError:
        pass
    return "body"


def extract(pptx_path):
    prs = Presentation(pptx_path)
    entries = []

    for slide_num, slide in enumerate(prs.slides, 1):
        for shape_idx, shape in enumerate(slide.shapes):
            kind = shape_kind(shape)

            if shape.has_text_frame:
                for para_idx, para in enumerate(shape.text_frame.paragraphs):
                    full = para.text.strip()
                    if not full:
                        continue
                    entries.append({
                        "slide": slide_num, "shape": shape_idx,
                        "run": para_idx, "kind": kind,
                        "zh": full, "en": "",
                    })

            elif shape.has_table:
                for row_idx, row in enumerate(shape.table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        text = cell.text.strip()
                        if not text:
                            continue
                        entries.append({
                            "slide": slide_num, "shape": shape_idx,
                            "run": len(shape.table.rows) * col_idx + row_idx,
                            "kind": "table", "zh": text, "en": "",
                        })

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                entries.append({
                    "slide": slide_num, "shape": -1, "run": 0,
                    "kind": "notes", "zh": notes, "en": "",
                })

    return entries


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: ./extract.py <input.pptx> <output.yaml>", file=sys.stderr)
        sys.exit(1)
    entries = extract(sys.argv[1])
    with open(sys.argv[2], "w") as f:
        yaml.dump(entries, f, allow_unicode=True, default_flow_style=False, sort_keys=False)
    print(f"Extracted {len(entries)} entries to {sys.argv[2]}")
